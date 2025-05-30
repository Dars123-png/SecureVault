from __future__ import print_function
import os.path
import base64
from datetime import datetime, timedelta, timezone
import io
import re

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

from email import message_from_bytes
from docx import Document
import openpyxl
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import csv
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import inch
from reportlab.lib import colors
import getpass

SCOPES = ['https://www.googleapis.com/auth/gmail.readonly']
# Sensitive Data Patterns
SENSITIVE_PATTERNS = {
    "SSN": r"\b\d{3}-\d{2}-\d{4}\b",
    "Credit Card": r"\b(?:4[0-9]{12}(?:[0-9]{3})?)\b",  # Visa
    "AWS Key": r"\bAKIA[0-9A-Z]{16}\b",
    "Password": r"\bpassword\s*[:=]\s*\S+",
    "API Key": r"\b(?:api|secret)[-_]?(key|token)?\s*[:=]\s*[A-Za-z0-9+/=._-]{10,}\b"
}

LOG_FILE = "email_monitor_log.txt"

# Declare tree as a global variable
tree = None

def log_violation(email_type, sender, label, snippet):
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(f"[{datetime.now()}] [{email_type}] From: {sender} | Type: {label} | Snippet: {snippet[:100]}\n")

def scan_text(text, sender, email_type, part):
    for label, pattern in SENSITIVE_PATTERNS.items():
        if re.search(pattern, text, re.IGNORECASE):
            print(f"[!] Sensitive data detected: {label} in {part} from {sender}")
            log_violation(email_type, sender, label, text)

def extract_text_from_attachment(part, msg_id, service):
    mime_type = part['mimeType']
    attachment_id = part['body']['attachmentId']
    data = service.users().messages().attachments().get(userId='me', messageId=msg_id, id=attachment_id).execute()
    file_data = base64.urlsafe_b64decode(data['data'].encode('UTF-8'))
    content = ""

    if 'docx' in mime_type:
        doc = Document(io.BytesIO(file_data))
        content = "\n".join(p.text for p in doc.paragraphs)
    elif 'pdf' in mime_type:
        from PyPDF2 import PdfReader  # Ensure you have PyPDF2 installed
        reader = PdfReader(io.BytesIO(file_data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except:
                print("[!] Encrypted PDF skipped.")
                return ""
        content = "\n".join(page.extract_text() or '' for page in reader.pages)
    elif 'sheet' in mime_type or 'excel' in mime_type:
        wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                content += " ".join(str(cell) for cell in row if cell) + "\n"
    return content

def get_gmail_service():
    creds = None
    if os.path.exists('token.json'):
        creds = Credentials.from_authorized_user_file('token.json', SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file('credentials.json', SCOPES)
            creds = flow.run_local_server(port=0)
        with open('token.json', 'w') as token:
            token.write(creds.to_json())
    return build('gmail', 'v1', credentials=creds)

def check_emails(service):
    print("[*] Checking emails for sensitive data...")

    ten_days_ago = (datetime.now(timezone.utc) - timedelta(days=10)).strftime("%Y/%m/%d")

    for label in ['INBOX', 'SENT']:
        print(f"\n[Checking {label} mails from last 10 days...]")
        query = f"after:{ten_days_ago}"
        next_page_token = None

        while True:
            results = service.users().messages().list(
                userId='me',
                labelIds=[label],
                q=query,
                pageToken=next_page_token
            ).execute()

            messages = results.get('messages', [])
            next_page_token = results.get('nextPageToken')

            for msg in messages:
                email = service.users().messages().get(userId='me', id=msg['id'], format='full').execute()
                payload = email.get("payload", {})
                headers = payload.get("headers", [])
                parts = payload.get("parts", [])

                sender = next((h['value'] for h in headers if h['name'] == 'From'), "Unknown")
                recipient = next((h['value'] for h in headers if h['name'] == 'To'), "Unknown")
                subject = next((h['value'] for h in headers if h['name'] == 'Subject'), "No Subject")
                date_str = next((h['value'] for h in headers if h['name'] == 'Date'), None)

                try:
                    date_obj = datetime.strptime(date_str, '%a, %d %b %Y %H:%M:%S %z')
                    formatted_date = date_obj.strftime('%Y-%m-%d %H:%M:%S %Z')
                except:
                    formatted_date = date_str or "Unknown Date"

                # Try scanning both plain and HTML parts
                scanned = False
                for part in parts:
                    mime_type = part.get("mimeType", "")
                    if mime_type in ["text/plain", "text/html"]:
                        data = part['body'].get("data")
                        if data:
                            body = base64.urlsafe_b64decode(data).decode("utf-8", errors="ignore")
                            if any(re.search(p, body, re.IGNORECASE) for p in SENSITIVE_PATTERNS.values()):
                                print(f"\n[+] From: {sender}")
                                if label == "SENT":
                                    print(f"    To: {recipient}")
                                print(f"    Subject: {subject}")
                                print(f"    Date: {formatted_date}")
                                print("[Body Preview]\n" + body[:500])
                                scan_text(body, sender, label, f"Email Body ({mime_type})")
                                scanned = True
                                break  # Stop once detected in body

                # Fallback for single-body emails
                if not parts and 'body' in payload:
                    data = payload['body'].get("data")
                    if data:
                        body = base64.urlsafe_b64decode(data).decode("utf-8", errors="ignore")
                        if any(re.search(p, body, re.IGNORECASE) for p in SENSITIVE_PATTERNS.values()):
                            print(f"\n[+] From: {sender}")
                            if label == "SENT":
                                print(f"    To: {recipient}")
                            print(f"    Subject: {subject}")
                            print(f"    Date: {formatted_date}")
                            print("[Body Preview]\n" + body[:300])
                            scan_text(body, sender, label, "Email Body (No Parts)")
                            scanned = True

                # Only check attachments if body was not flagged
                if not scanned:
                    for part in parts:
                        if part.get("filename") and 'attachmentId' in part.get('body', {}):
                            try:
                                content = extract_text_from_attachment(part, msg['id'], service)
                                if any(re.search(p, content, re.IGNORECASE) for p in SENSITIVE_PATTERNS.values()):
                                    print(f"\n[+] From: {sender}")
                                    if label == "SENT":
                                        print(f"    To: {recipient}")
                                    print(f"    Subject: {subject}")
                                    print(f"    Date: {formatted_date}")
                                    print(f"[+] Scanning attachment: {part['filename']}")
                                    scan_text(content, sender, label, f"Attachment: {part['filename']}")
                            except Exception as e:
                                print(f"[!] Failed to scan attachment: {e}")

            if not next_page_token:
                break
    print("\n[*] Scan complete.")

def load_logs(tree):
    tree.delete(*tree.get_children())
    if not os.path.exists(LOG_FILE):
        return
    with open(LOG_FILE, "r", encoding="utf-8") as f:
        for line in f:
            try:
                timestamp = line.split("]")[0].strip("[")
                rest = line.split("]")[2]
                sender = rest.split("|")[0].split(":")[1].strip()
                label = rest.split("|")[1].split(":")[1].strip()
                tree.insert("", "end", values=(sender, timestamp, label))
            except:
                continue

def export_csv():
    if not os.path.exists(LOG_FILE):
        messagebox.showerror("Error", "Log file not found.")
        return
    file_path = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV files", "*.csv")])
    if file_path:
        with open(LOG_FILE, "r", encoding="utf-8") as f, open(file_path, "w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            writer.writerow(["Sender", "Timestamp", "Sensitive Type"])
            for line in f:
                try:
                    timestamp = line.split("]")[0].strip("[")
                    rest = line.split("]")[2]
                    sender = rest.split("|")[0].split(":")[1].strip()
                    label = rest.split("|")[1].split(":")[1].strip()
                    writer.writerow([sender, timestamp, label])
                except:
                    continue
        messagebox.showinfo("Success", "Logs exported to CSV successfully.")

def export_pdf(tree):
    if not os.path.exists(LOG_FILE):
        messagebox.showerror("Error", "Log file not found.")
        return

    file_path = filedialog.asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")])
    if not file_path:
        return

    try:
        # Handle both Treeview widget and list of data rows
        if hasattr(tree, 'get_children'):
            rows = [tree.item(item)['values'] for item in tree.get_children()]
        elif isinstance(tree, list):
            rows = tree
        else:
            messagebox.showerror("Error", "Invalid data format passed to export_pdf().")
            return

        doc = SimpleDocTemplate(file_path, pagesize=A4,
                                rightMargin=40, leftMargin=40,
                                topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("TitleStyle", parent=styles['Title'], alignment=TA_CENTER,
                                     fontSize=16, leading=20, spaceAfter=10, spaceBefore=10, allCaps=True)
        cell_style = ParagraphStyle("CellStyle", fontSize=9, leading=12)

        Story = [Paragraph("SENSITIVE EMAIL DETECTION REPORT", title_style), Spacer(1, 0.2 * inch)]

        if rows:
            # Prepare header
            table_data = [["Sender", "Timestamp", "Sensitive Data"]]

            # Wrap long text in Paragraphs
            for row in rows:
                formatted_row = [
                    Paragraph(str(row[0]), cell_style),
                    Paragraph(str(row[1]), cell_style),
                    Paragraph(str(row[2]), cell_style)
                ]
                table_data.append(formatted_row)

            table = Table(table_data, colWidths=[200, 150, 100])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                ('BACKGROUND', (0, 1), (-1, -1), colors.whitesmoke),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ]))
            Story.append(table)
        else:
            Story.append(Paragraph("No dashboard data available.", styles['Normal']))

        doc.build(Story)
        messagebox.showinfo("Success", "PDF exported successfully.")
    except Exception as e:
        messagebox.showerror("Error", f"Failed to export PDF: {str(e)}")

def start_monitoring(tree):
    service = get_gmail_service()
    check_emails(service)
    load_logs(tree)
    messagebox.showinfo("Monitoring", "Scan complete and log updated.")

def clear_logs():
    global tree  # Declare tree as global
    if messagebox.askyesno("Confirm", "Are you sure you want to clear all logs?"):
        open(LOG_FILE, 'w').close()  # Clear the log file
        load_logs(tree)  # Refresh the log display

# Tkinter GUI
def create_dashboard():
    global tree  # Declare tree as global
    root = tk.Tk()
    root.title("Gmail Sensitive Data Monitor")
    root.geometry("800x500")

    # Buttons
    btn_frame = tk.Frame(root)
    btn_frame.pack(pady=10)

    tree = ttk.Treeview(root, columns=("Sender", "Timestamp", "Sensitive Data"), show="headings")
    tree.heading("Sender", text="Sender")
    tree.heading("Timestamp", text="Timestamp")
    tree.heading("Sensitive Data", text="Sensitive Data")
    tree.column("Sender", width=250)
    tree.column("Timestamp", width=200)
    tree.column("Sensitive Data", width=200)
    tree.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

    # Create buttons with unique column indices
    tk.Button(btn_frame, text="Start Monitoring", command=lambda: start_monitoring(tree), width=18).grid(row=0, column=0, padx=10)
    tk.Button(btn_frame, text="Export to CSV", command=export_csv, width=18).grid(row=0, column=1, padx=10)
    tk.Button(btn_frame, text="Clear Logs", command=clear_logs, width=15).grid(row=0, column=2, padx=10)
    tk.Button(
        btn_frame,
        text="Export to PDF",
        command=lambda: export_pdf([tree.item(i)['values'] for i in tree.get_children()]),
        width=18
    ).grid(row=0, column=3, padx=10)

    load_logs(tree)
    root.mainloop()

# Launch the GUI
if __name__ == "__main__":
    create_dashboard()
